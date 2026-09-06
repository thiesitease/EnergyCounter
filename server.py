#!/usr/bin/env python3
"""Zählerbuch SB85 – kleiner Server für die selbst gehostete Variante.

Aufgaben:
  * liefert die Web-Seite (static/index.html)
  * speichert Ablesungen und Zähler-Konfiguration als JSON-Dateien unter data/
  * Foto-Erkennung: schickt ein Zählerfoto an die Anthropic-API (Claude) und gibt
    Zähler, Zählerstand und Sicherheit als JSON zurück
  * Wetterdaten: holt Tagesmitteltemperaturen für den Standort von Open-Meteo
    (kostenlos, ohne Schlüssel) und cached sie in data/weather.json
  * Excel-Export aller Ablesungen inkl. Temperaturen

Umgebungsvariablen:
  ANTHROPIC_API_KEY  Schlüssel von console.anthropic.com (nötig für Fotos)
  ZB_PASSWORD        Passwort für die Seite (leer = kein Schutz, nur im LAN sinnvoll)
  ZB_MODEL           Modell für die Erkennung (Standard: claude-opus-5)
  ZB_DATA            Datenverzeichnis (Standard: ./data)
  PORT               Port (Standard: 8080)
"""
import base64
import datetime as dt
import hmac
import io
import json
import os
import re
import threading
import time
import urllib.parse
import urllib.request
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

ROOT = Path(__file__).resolve().parent
DATA = Path(os.environ.get("ZB_DATA", ROOT / "data"))
STATIC = ROOT / "static"
PORT = int(os.environ.get("PORT", "8080"))
PASSWORD = os.environ.get("ZB_PASSWORD", "")
MODEL = os.environ.get("ZB_MODEL", "claude-opus-5")
LOCK = threading.Lock()
DAY = 86400000

# --------------------------------------------------------------------------- storage

def _load(name, default):
    p = DATA / name
    if not p.exists():
        return default
    with open(p, encoding="utf-8") as f:
        return json.load(f)


def _save(name, obj):
    DATA.mkdir(parents=True, exist_ok=True)
    tmp = DATA / (name + ".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)
    os.replace(tmp, DATA / name)


def readings():
    return sorted(_load("readings.json", []), key=lambda r: r.get("t", 0))


def config():
    return _load("config.json", {})


def weather():
    return _load("weather.json", {"daily": {}})


# --------------------------------------------------------------------------- weather (Open-Meteo)

def _get_json(url, timeout=30):
    req = urllib.request.Request(url, headers={"User-Agent": "zaehlerbuch/1.0"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))


def _fetch_daily(base, lat, lon, params):
    q = {"latitude": lat, "longitude": lon, "daily": "temperature_2m_mean,temperature_2m_min,temperature_2m_max",
         "timezone": "Europe/Berlin", **params}
    data = _get_json(base + "?" + urllib.parse.urlencode(q))
    d = data.get("daily", {})
    out = {}
    for i, day in enumerate(d.get("time", [])):
        mean = d["temperature_2m_mean"][i]
        if mean is None:
            continue
        out[day] = {"mean": mean, "min": d["temperature_2m_min"][i], "max": d["temperature_2m_max"][i]}
    return out


_weather_busy = threading.Event()


def update_weather(force=False):
    """Fehlende Tage nachladen. Läuft im Hintergrund; gibt sofort zurück, wenn schon aktiv."""
    cfg = config()
    loc = cfg.get("location") or {}
    if not loc.get("lat") or not loc.get("lon"):
        return False
    if _weather_busy.is_set():
        return True
    w = weather()
    same_loc = abs(float(w.get("lat", 0)) - float(loc["lat"])) < 1e-4 and abs(float(w.get("lon", 0)) - float(loc["lon"])) < 1e-4
    if not force and same_loc and time.time() - w.get("updated", 0) < 6 * 3600:
        return False

    def work():
        _weather_busy.set()
        try:
            daily = dict(w.get("daily", {})) if same_loc else {}
            rs = readings()
            first = dt.date.fromtimestamp(rs[0]["t"] / 1000) if rs else dt.date.today() - dt.timedelta(days=365)
            first = min(first, dt.date.today() - dt.timedelta(days=400))
            today = dt.date.today()
            archive_end = today - dt.timedelta(days=7)
            # Archiv in Blöcken von 5 Jahren, nur was fehlt
            start = first
            if daily:
                have = sorted(daily)
                start = max(first, dt.date.fromisoformat(have[-1]) - dt.timedelta(days=10))
            while start <= archive_end:
                end = min(archive_end, start + dt.timedelta(days=5 * 365))
                daily.update(_fetch_daily("https://archive-api.open-meteo.com/v1/archive", loc["lat"], loc["lon"],
                                          {"start_date": start.isoformat(), "end_date": end.isoformat()}))
                start = end + dt.timedelta(days=1)
            # die letzten zwei Wochen aus dem Vorhersage-Endpunkt (Archiv hängt einige Tage hinterher)
            recent = _fetch_daily("https://api.open-meteo.com/v1/forecast", loc["lat"], loc["lon"],
                                  {"past_days": 14, "forecast_days": 1})
            for day, v in recent.items():
                if day <= today.isoformat():
                    daily[day] = v
            with LOCK:
                _save("weather.json", {"lat": loc["lat"], "lon": loc["lon"], "name": loc.get("name", ""),
                                       "updated": time.time(), "daily": dict(sorted(daily.items()))})
        except Exception as e:  # noqa: BLE001
            print("Wetter-Update fehlgeschlagen:", e)
        finally:
            _weather_busy.clear()

    threading.Thread(target=work, daemon=True).start()
    return True


# --------------------------------------------------------------------------- photo recognition (Anthropic)

RESULT_SCHEMA = {
    "type": "object",
    "properties": {
        "meter": {"type": "string", "description": "key of the known meter, or 'unknown'"},
        "id_text": {"type": "string", "description": "serial number as printed, or empty"},
        "reading": {"type": ["number", "null"], "description": "counter value as decimal number in the meter's unit"},
        "confidence": {"type": "number", "minimum": 0, "maximum": 1},
        "comment": {"type": "string", "description": "short German note, or empty"},
    },
    "required": ["meter", "id_text", "reading", "confidence", "comment"],
    "additionalProperties": False,
}


def recognize(image_b64, media_type, prompt):
    import anthropic  # erst hier, damit der Server auch ohne Schlüssel startet

    if not os.environ.get("ANTHROPIC_API_KEY") and not os.environ.get("ANTHROPIC_AUTH_TOKEN"):
        return {"error": "no_api_key", "message": "ANTHROPIC_API_KEY ist auf dem Server nicht gesetzt."}
    client = anthropic.Anthropic()
    content = [
        {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": image_b64}},
        {"type": "text", "text": prompt},
    ]
    kwargs = dict(
        model=MODEL,
        max_tokens=4000,
        messages=[{"role": "user", "content": content}],
        output_config={"format": {"type": "json_schema", "schema": RESULT_SCHEMA}},
    )
    try:
        try:
            # Server-seitiger Fallback: wird eine Anfrage vom Sicherheitsfilter abgelehnt,
            # läuft sie automatisch auf einem anderen Modell weiter.
            resp = client.beta.messages.create(betas=["server-side-fallback-2026-07-01"],
                                               extra_body={"fallbacks": "default"}, **kwargs)
        except anthropic.BadRequestError as e:
            if "fallback" not in str(e).lower():
                raise
            resp = client.beta.messages.create(**kwargs)
    except anthropic.AuthenticationError:
        return {"error": "auth", "message": "API-Schlüssel ungültig."}
    except anthropic.RateLimitError:
        return {"error": "rate_limited", "message": "Zu viele Anfragen, bitte kurz warten."}
    except anthropic.APIStatusError as e:
        return {"error": "api", "message": f"API-Fehler {e.status_code}: {e.message}"}
    except anthropic.APIConnectionError:
        return {"error": "network", "message": "Keine Verbindung zur Anthropic-API."}
    if resp.stop_reason == "refusal":
        return {"error": "refused", "message": "Claude hat dieses Bild abgelehnt."}
    text = next((b.text for b in resp.content if b.type == "text"), "")
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", text, re.S)
        if not m:
            return {"error": "invalid_json", "message": "Antwort nicht lesbar.", "text": text}
        data = json.loads(m.group(0))
    u = resp.usage
    data["usage"] = {"input_tokens": u.input_tokens, "output_tokens": u.output_tokens, "model": resp.model}
    return data


# --------------------------------------------------------------------------- analysis helpers (für Export)

def factor_of(m):
    f = m.get("factors")
    if not f:
        return 1.0
    return float(f.get("zz") or 1) * float(f.get("bw") or 1)


def price_segments(m):
    segs = []
    for p in m.get("prices", []):
        try:
            t = dt.datetime.fromisoformat(p["from"]).timestamp() * 1000
        except (KeyError, ValueError):
            continue
        segs.append((t, float(p.get("price") or 0)))
    return sorted(segs)


def cost_between(m, a, b, amount):
    segs = price_segments(m)
    if not segs or b <= a:
        return 0.0
    cost = 0.0
    for i, (t, price) in enumerate(segs):
        s = max(a, t)
        e = min(b, segs[i + 1][0] if i + 1 < len(segs) else float("inf"))
        if e > s:
            cost += amount * (e - s) / (b - a) * price
    return cost


def series(rs, m):
    """Punkte mit kumuliertem Stand (Zählerwechsel berücksichtigt) und Intervallwerten."""
    key = m["key"]
    f = factor_of(m)
    pts, addon, prev = [], 0.0, None
    for r in rs:
        raw = (r.get("v") or {}).get(key)
        if not isinstance(raw, (int, float)):
            continue
        chg = bool((r.get("chg") or {}).get(key))
        if chg and prev is not None:
            addon = prev
        cum_raw = raw + addon
        pts.append({"t": r["t"], "raw": raw, "cum": cum_raw * f, "chg": chg, "note": (r.get("n") or {}).get(key, "")})
        prev = cum_raw
    for i in range(1, len(pts)):
        a, b = pts[i - 1], pts[i]
        b["days"] = (b["t"] - a["t"]) / DAY
        b["amount"] = b["cum"] - a["cum"]
        b["rate"] = b["amount"] / b["days"] if b["days"] > 0.02 else None
        b["cost"] = cost_between(m, a["t"], b["t"], b["amount"])
    return pts


def temp_stats(daily, a, b):
    """Mittlere Außentemperatur und Gradtagzahl (20/15) für [a, b] in ms."""
    da = dt.date.fromtimestamp(a / 1000)
    db = dt.date.fromtimestamp(b / 1000)
    temps = []
    d = da
    while d <= db:
        v = daily.get(d.isoformat())
        if v is not None:
            temps.append(v["mean"] if isinstance(v, dict) else v)
        d += dt.timedelta(days=1)
    if not temps:
        return None, None, 0
    mean = sum(temps) / len(temps)
    gtz = sum(20 - t for t in temps if t < 15)
    return mean, gtz, len(temps)


def build_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    rs = readings()
    cfg = config()
    ms = cfg.get("meters") or [{"key": k, "name": k, "unit": ""} for k in
                               sorted({k for r in rs for k in (r.get("v") or {})})]
    daily = weather().get("daily", {})
    wb = Workbook()
    bold = Font(name="Arial", bold=True)
    normal = Font(name="Arial")

    def ws_header(ws, cols):
        ws.append(cols)
        for c in ws[1]:
            c.font = bold
        ws.freeze_panes = "A2"

    # Blatt 1: Zählerstände (wie die Übersichtstabelle in Excel)
    ws = wb.active
    ws.title = "Zählerstände"
    cols = ["Datum"] + [f"{m['name']} ({m.get('unit','')})" for m in ms] + ["Ø Außentemp. seit letzter Ablesung °C", "Gradtage (20/15) seit letzter Ablesung"] \
        + [f"Wechsel {m['name']}" for m in ms] + [f"Notiz {m['name']}" for m in ms] + ["Quelle"]
    ws_header(ws, cols)
    prev_t = None
    for r in rs:
        t = dt.datetime.fromtimestamp(r["t"] / 1000)
        mean, gtz, n = (None, None, 0)
        if prev_t is not None:
            mean, gtz, n = temp_stats(daily, prev_t, r["t"])
        row = [t] + [(r.get("v") or {}).get(m["key"]) for m in ms] + [round(mean, 1) if mean is not None else None, round(gtz, 1) if gtz is not None else None] \
            + ["x" if (r.get("chg") or {}).get(m["key"]) else None for m in ms] + [(r.get("n") or {}).get(m["key"]) for m in ms] + [r.get("src", "")]
        ws.append(row)
        ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY HH:MM"
        prev_t = r["t"]
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = normal
    ws.column_dimensions["A"].width = 18
    for i in range(2, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16

    # Blatt 2: Verbrauch je Abschnitt und Zähler
    ws2 = wb.create_sheet("Verbrauch")
    ws_header(ws2, ["Zähler", "Von", "Bis", "Tage", "Stand von", "Stand bis", "Verbrauch", "Einheit", "Verbrauch/Tag", "Kosten €", "Kosten/Tag €",
                    "Ø Außentemp. °C", "Gradtage (20/15)", "Verbrauch je Gradtag", "Zählerwechsel", "Notiz"])
    for m in ms:
        pts = series(rs, m)
        unit = m.get("outUnit") or m.get("unit", "")
        for i in range(1, len(pts)):
            a, b = pts[i - 1], pts[i]
            mean, gtz, n = temp_stats(daily, a["t"], b["t"])
            ws2.append([m["name"], dt.datetime.fromtimestamp(a["t"] / 1000), dt.datetime.fromtimestamp(b["t"] / 1000), round(b["days"], 2),
                        a["raw"], b["raw"], round(b["amount"], 3), unit, round(b["rate"], 3) if b["rate"] is not None else None,
                        round(b["cost"], 2), round(b["cost"] / b["days"], 3) if b["days"] > 0.02 else None,
                        round(mean, 1) if mean is not None else None, round(gtz, 1) if gtz is not None else None,
                        round(b["amount"] / gtz, 3) if gtz else None, "x" if b["chg"] else None, b["note"] or None])
            ws2.cell(ws2.max_row, 2).number_format = "DD.MM.YYYY HH:MM"
            ws2.cell(ws2.max_row, 3).number_format = "DD.MM.YYYY HH:MM"
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.font = normal
    for i, w in enumerate([18, 17, 17, 7, 12, 12, 12, 8, 13, 10, 11, 14, 14, 16, 12, 40], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # Blatt 3: Temperaturen (Tageswerte)
    ws3 = wb.create_sheet("Temperaturen")
    ws_header(ws3, ["Datum", "Mittel °C", "Min °C", "Max °C"])
    for day, v in sorted(daily.items()):
        ws3.append([dt.date.fromisoformat(day), v.get("mean"), v.get("min"), v.get("max")])
        ws3.cell(ws3.max_row, 1).number_format = "DD.MM.YYYY"
    for row in ws3.iter_rows(min_row=2):
        for c in row:
            c.font = normal
    ws3.column_dimensions["A"].width = 12
    loc = cfg.get("location") or {}
    ws3["F1"] = "Quelle: Open-Meteo (open-meteo.com), Standort " + str(loc.get("name", "")) + f" ({loc.get('lat','')}, {loc.get('lon','')})"
    ws3["F1"].font = normal

    # Blatt 4: Zähler und Preise
    ws4 = wb.create_sheet("Zähler")
    ws_header(ws4, ["Zähler", "Schlüssel", "Einheit", "Ergebnis-Einheit", "Faktor", "IDs", "in Summe", "Preis gültig ab", "Preis €"])
    for m in ms:
        prices = m.get("prices") or [{}]
        for p in prices:
            ws4.append([m["name"], m["key"], m.get("unit", ""), m.get("outUnit") or m.get("unit", ""), factor_of(m), ", ".join(m.get("ids", [])),
                        "x" if m.get("inTotal") else None, p.get("from"), p.get("price")])
    for row in ws4.iter_rows(min_row=2):
        for c in row:
            c.font = normal
    for i, w in enumerate([20, 10, 8, 14, 8, 34, 8, 14, 9], start=1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- HTTP

class Handler(BaseHTTPRequestHandler):
    server_version = "Zaehlerbuch/1.0"

    def log_message(self, fmt, *args):  # kürzeres Log
        print(dt.datetime.now().strftime("%H:%M:%S"), self.address_string(), fmt % args)

    # ---- helpers
    def _send(self, code, body=b"", ctype="application/json; charset=utf-8", extra=None):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        if self.command != "HEAD":
            self.wfile.write(body)

    def _json(self, obj, code=200):
        self._send(code, json.dumps(obj, ensure_ascii=False).encode("utf-8"))

    def _body(self):
        n = int(self.headers.get("Content-Length") or 0)
        if n > 25 * 1024 * 1024:
            raise ValueError("body too large")
        raw = self.rfile.read(n) if n else b""
        return json.loads(raw.decode("utf-8")) if raw else {}

    def _authed(self):
        if not PASSWORD:
            return True
        auth = self.headers.get("Authorization", "")
        if auth.startswith("Bearer ") and hmac.compare_digest(auth[7:], PASSWORD):
            return True
        cookie = self.headers.get("Cookie", "")
        m = re.search(r"(?:^|;\s*)zb_auth=([^;]+)", cookie)
        if m and hmac.compare_digest(urllib.parse.unquote(m.group(1)), PASSWORD):
            return True
        return False

    # ---- routes
    def do_GET(self):
        path, _, query = self.path.partition("?")
        qs = urllib.parse.parse_qs(query)
        if path in ("/", "/index.html"):
            return self._static("index.html", "text/html; charset=utf-8")
        if path == "/manifest.webmanifest":
            return self._static("manifest.webmanifest", "application/manifest+json")
        if path == "/api/health":
            return self._json({"ok": True, "readings": len(readings()), "model": MODEL,
                               "api_key": bool(os.environ.get("ANTHROPIC_API_KEY") or os.environ.get("ANTHROPIC_AUTH_TOKEN")),
                               "password": bool(PASSWORD)})
        if not path.startswith("/api/"):
            return self._send(404, b"not found", "text/plain")
        if not self._authed():
            return self._json({"error": "unauthorized"}, 401)
        if path == "/api/state":
            refreshing = update_weather()
            w = weather()
            return self._json({"readings": readings(), "config": config(),
                               "weather": {"daily": {d: v["mean"] for d, v in w.get("daily", {}).items()},
                                           "updated": w.get("updated"), "refreshing": refreshing or _weather_busy.is_set(),
                                           "name": w.get("name", "")},
                               "features": {"photo": bool(os.environ.get("ANTHROPIC_API_KEY") or os.environ.get("ANTHROPIC_AUTH_TOKEN")), "model": MODEL}})
        if path == "/api/weather":
            refreshing = update_weather(force="refresh" in qs)
            w = weather()
            return self._json({"daily": {d: v["mean"] for d, v in w.get("daily", {}).items()}, "updated": w.get("updated"),
                               "refreshing": refreshing or _weather_busy.is_set(), "name": w.get("name", "")})
        if path == "/api/geocode":
            q = (qs.get("q") or [""])[0].strip()
            if not q:
                return self._json([])
            try:
                data = _get_json("https://geocoding-api.open-meteo.com/v1/search?" + urllib.parse.urlencode({"name": q, "count": 6, "language": "de", "format": "json"}))
            except Exception as e:  # noqa: BLE001
                return self._json({"error": "geocode", "message": str(e)}, 502)
            return self._json([{"name": r.get("name"), "admin": r.get("admin1", ""), "country": r.get("country", ""),
                                "lat": r.get("latitude"), "lon": r.get("longitude")} for r in data.get("results", [])])
        if path == "/api/export.xlsx":
            try:
                body = build_xlsx()
            except Exception as e:  # noqa: BLE001
                return self._json({"error": "export", "message": str(e)}, 500)
            name = "Zaehlerbuch_" + dt.date.today().isoformat() + ".xlsx"
            return self._send(200, body, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                              {"Content-Disposition": f'attachment; filename="{name}"'})
        return self._json({"error": "not found"}, 404)

    def do_HEAD(self):
        self.do_GET()

    def do_POST(self):
        path = self.path.partition("?")[0]
        if path == "/api/login":
            try:
                pw = str(self._body().get("password", ""))
            except Exception:  # noqa: BLE001
                pw = ""
            if not PASSWORD or hmac.compare_digest(pw, PASSWORD):
                return self._json({"ok": True})
            return self._json({"ok": False}, 401)
        if not self._authed():
            return self._json({"error": "unauthorized"}, 401)
        if path == "/api/recognize":
            try:
                b = self._body()
                img = str(b.get("image", ""))
                if img.startswith("data:"):
                    img = img.split(",", 1)[1]
                media = str(b.get("media_type") or "image/jpeg")
                if media not in ("image/jpeg", "image/png", "image/webp", "image/gif"):
                    media = "image/jpeg"
                base64.b64decode(img[:64] + "=" * (-len(img[:64]) % 4))  # grober Formatcheck
                result = recognize(img, media, str(b.get("prompt", "")))
            except Exception as e:  # noqa: BLE001
                return self._json({"error": "server", "message": str(e)}, 500)
            return self._json(result, 200 if "error" not in result else 502)
        return self._json({"error": "not found"}, 404)

    def do_PUT(self):
        path = self.path.partition("?")[0]
        if not self._authed():
            return self._json({"error": "unauthorized"}, 401)
        try:
            body = self._body()
        except Exception as e:  # noqa: BLE001
            return self._json({"error": "bad json", "message": str(e)}, 400)
        m = re.fullmatch(r"/api/readings/([A-Za-z0-9_\-]{1,64})", path)
        if m:
            rid = m.group(1)
            if not isinstance(body.get("t"), (int, float)) or not isinstance(body.get("v"), dict):
                return self._json({"error": "invalid reading"}, 400)
            doc = {"id": rid, "t": int(body["t"]), "v": {k: float(v) for k, v in body["v"].items() if isinstance(v, (int, float))},
                   "src": str(body.get("src") or "manual")}
            if isinstance(body.get("n"), dict) and body["n"]:
                doc["n"] = {k: str(v)[:500] for k, v in body["n"].items() if v}
            if isinstance(body.get("chg"), dict) and body["chg"]:
                doc["chg"] = {k: True for k, v in body["chg"].items() if v}
            for k in ("created", "edited"):
                if isinstance(body.get(k), (int, float)):
                    doc[k] = int(body[k])
            with LOCK:
                rs = [r for r in _load("readings.json", []) if r.get("id") != rid]
                rs.append(doc)
                _save("readings.json", sorted(rs, key=lambda r: r.get("t", 0)))
            return self._json({"ok": True})
        if path == "/api/config":
            if not isinstance(body.get("meters"), list):
                return self._json({"error": "invalid config"}, 400)
            with LOCK:
                _save("config.json", body)
            update_weather()
            return self._json({"ok": True})
        return self._json({"error": "not found"}, 404)

    def do_DELETE(self):
        path = self.path.partition("?")[0]
        if not self._authed():
            return self._json({"error": "unauthorized"}, 401)
        m = re.fullmatch(r"/api/readings/([A-Za-z0-9_\-]{1,64})", path)
        if not m:
            return self._json({"error": "not found"}, 404)
        with LOCK:
            rs = [r for r in _load("readings.json", []) if r.get("id") != m.group(1)]
            _save("readings.json", rs)
        return self._json({"ok": True})

    def _static(self, name, ctype):
        p = STATIC / name
        if not p.exists():
            return self._send(404, b"missing " + name.encode(), "text/plain")
        self._send(200, p.read_bytes(), ctype)


def main():
    DATA.mkdir(parents=True, exist_ok=True)
    print(f"Zählerbuch läuft auf http://0.0.0.0:{PORT}  (Daten: {DATA}, Modell: {MODEL}, "
          f"Passwort: {'ja' if PASSWORD else 'nein'}, API-Key: {'ja' if os.environ.get('ANTHROPIC_API_KEY') else 'nein'})")
    update_weather()
    ThreadingHTTPServer(("0.0.0.0", PORT), Handler).serve_forever()


if __name__ == "__main__":
    main()
